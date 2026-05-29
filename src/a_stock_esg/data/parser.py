"""
文档解析器
支持PDF、HTML、Word等格式的ESG报告解析
"""

from dataclasses import dataclass
from typing import Optional, Dict, List
from pathlib import Path


@dataclass
class ParsedDocument:
    """解析后的文档"""
    file_path: str
    file_type: str
    content: str
    metadata: Dict = None
    
    def __post_init__(self):
        if self.metadata is None:
            self.metadata = {}


class DocumentParser:
    """
    文档解析器
    
    支持PDF、HTML、Word等格式的ESG报告解析
    """
    
    def __init__(self):
        """初始化解析器"""
        self.supported_formats = [".pdf", ".html", ".docx", ".txt"]
    
    def parse(self, file_path: str) -> ParsedDocument:
        """
        解析文档
        
        Args:
            file_path: 文件路径
            
        Returns:
            ParsedDocument: 解析后的文档
        """
        path = Path(file_path)
        
        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")
        
        if path.suffix.lower() not in self.supported_formats:
            raise ValueError(f"不支持的文件格式: {path.suffix}")
        
        if path.suffix.lower() == ".pdf":
            return self._parse_pdf(file_path)
        elif path.suffix.lower() == ".html":
            return self._parse_html(file_path)
        elif path.suffix.lower() == ".docx":
            return self._parse_docx(file_path)
        elif path.suffix.lower() == ".txt":
            return self._parse_txt(file_path)
        else:
            raise ValueError(f"未知的文件格式: {path.suffix}")
    
    def _parse_pdf(self, file_path: str) -> ParsedDocument:
        """解析PDF文件"""
        try:
            import pdfplumber
            
            content = []
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        content.append(text)
            
            return ParsedDocument(
                file_path=file_path,
                file_type="pdf",
                content="\n".join(content),
                metadata={"pages": len(content)},
            )
        except ImportError:
            # 备用方案：使用PyPDF2
            from PyPDF2 import PdfReader
            
            reader = PdfReader(file_path)
            content = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    content.append(text)
            
            return ParsedDocument(
                file_path=file_path,
                file_type="pdf",
                content="\n".join(content),
                metadata={"pages": len(reader.pages)},
            )
    
    def _parse_html(self, file_path: str) -> ParsedDocument:
        """解析HTML文件"""
        from bs4 import BeautifulSoup
        
        with open(file_path, "r", encoding="utf-8") as f:
            html_content = f.read()
        
        soup = BeautifulSoup(html_content, "lxml")
        
        # 移除脚本和样式
        for script in soup(["script", "style"]):
            script.decompose()
        
        text = soup.get_text(separator="\n", strip=True)
        
        # 提取标题
        title = soup.find("title")
        title_text = title.string if title else ""
        
        return ParsedDocument(
            file_path=file_path,
            file_type="html",
            content=text,
            metadata={"title": title_text},
        )
    
    def _parse_docx(self, file_path: str) -> ParsedDocument:
        """解析Word文档"""
        from docx import Document
        
        doc = Document(file_path)
        content = []
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                content.append(paragraph.text)
        
        # 提取表格内容
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells]
                content.append(" | ".join(row_text))
        
        return ParsedDocument(
            file_path=file_path,
            file_type="docx",
            content="\n".join(content),
            metadata={"paragraphs": len(doc.paragraphs)},
        )
    
    def _parse_txt(self, file_path: str) -> ParsedDocument:
        """解析纯文本文件"""
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()
        
        return ParsedDocument(
            file_path=file_path,
            file_type="txt",
            content=content,
        )
    
    def batch_parse(self, file_paths: List[str]) -> List[ParsedDocument]:
        """批量解析文档"""
        results = []
        
        for file_path in file_paths:
            try:
                doc = self.parse(file_path)
                results.append(doc)
            except Exception as e:
                print(f"解析文件失败 {file_path}: {e}")
        
        return results
